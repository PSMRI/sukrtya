"""Compact dump of the Form column in `SukrtyaQuestion.filled.xlsx` so we can verify the
per-row mapping at a glance.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    path = Path("SukrtyaQuestion.filled.xlsx")
    wb = load_workbook(filename=str(path), data_only=True)
    ws = wb["Question Master"]

    header_row = -1
    cols: dict[str, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=10):
        for cell in row:
            v = (str(cell.value).strip().lower().replace(" ", "") if cell.value else "")
            if v in ("faqid", "faquestionid", "qid"):
                header_row = cell.row
        if header_row > 0:
            for cell in ws[header_row]:
                v = (str(cell.value).strip().lower().replace(" ", "") if cell.value else "")
                if v in ("formcode", "form", "formname", "section", "sectioncode") and "form" not in cols:
                    cols["form"] = cell.column
                elif v in ("faqid", "faquestionid", "qid") and "qid" not in cols:
                    cols["qid"] = cell.column
                elif v in ("faquestionen", "questionen", "question", "faquestion") and "qen" not in cols:
                    cols["qen"] = cell.column
                elif v in ("faanswertype", "answertype", "type") and "atype" not in cols:
                    cols["atype"] = cell.column
            break

    print(f"file:        {path.resolve()}")
    print(f"header row:  {header_row}")
    print(f"Form column: {cols.get('form')}")
    print()
    print(f"{'Row':>3}  {'qID':>4}  {'Form':<8}  {'Type':<14}  Label")
    print("-" * 90)
    last_form = ""
    for r in range(header_row + 1, ws.max_row + 1):
        qid = ws.cell(row=r, column=cols["qid"]).value
        form = ws.cell(row=r, column=cols["form"]).value or ""
        qen = ws.cell(row=r, column=cols["qen"]).value or ""
        atype = ws.cell(row=r, column=cols["atype"]).value or ""
        if qid is None and not qen:
            continue
        marker = "<-- NEW SECTION" if str(form) != last_form and form else ""
        last_form = str(form)
        print(f"{r:>3}  {str(qid):>4}  {str(form):<8}  {str(atype):<14}  {str(qen)[:55]}  {marker}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
