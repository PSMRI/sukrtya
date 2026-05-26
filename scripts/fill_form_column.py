"""Generate `SukrtyaQuestion.filled.xlsx` from `SukrtyaQuestion.xlsx` with the Form column
explicitly populated on every question row.

Mirrors the FormExcelImportService section-detection logic so the output is exactly what
the backend would group rows into. After reviewing the file you can override any cell in the
Form column (e.g. switch some ANC4 rows to DELIVERY / PNC1 / PNC2) and re-upload.

Usage: `python scripts/fill_form_column.py`
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


SRC = Path("SukrtyaQuestion.xlsx")
DST = Path("SukrtyaQuestion.filled.xlsx")

FORM_CODE_ALIASES = {
    "BASIC": "BASIC", "BASICDETAILS": "BASIC", "BASICDETAIL": "BASIC", "PREGNANTWOMAN": "BASIC",
    "ANC1": "ANC1", "ANC2": "ANC2", "ANC3": "ANC3", "ANC4": "ANC4",
    "FORM1": "BASIC",
    "1STANC": "ANC1", "2NDANC": "ANC2", "3RDANC": "ANC3", "4THANC": "ANC4",
    "FIRSTANC": "ANC1", "SECONDANC": "ANC2", "THIRDANC": "ANC3", "FOURTHANC": "ANC4",
}

ANC1_RE = re.compile(r"\b(1st\s*anc|first\s*anc|anc\s*1\b|of\s*1\s*anc)\b", re.I)
ANC2_RE = re.compile(r"\b(2nd\s*anc|second\s*anc|anc\s*2\b|of\s*2\s*anc)\b", re.I)
ANC3_RE = re.compile(r"\b(3rd\s*anc|third\s*anc|anc\s*3\b|of\s*3\s*anc)\b", re.I)
ANC4_RE = re.compile(r"\b(4th\s*anc|fourth\s*anc|anc\s*4\b|of\s*4\s*anc)\b", re.I)

DEFAULT_FORM = "BASIC"
HIGHLIGHT = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")


def header_key(s: object) -> str:
    if s is None:
        return ""
    return str(s).strip().lower().replace(" ", "")


def is_blank(s: object) -> bool:
    if s is None:
        return True
    t = str(s).strip()
    return t == "" or t.lower() == "null" or t == "-"


def normalise(raw: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", raw.strip().upper())


def infer(label: Optional[str]) -> Optional[str]:
    if not label:
        return None
    lower = label.lower()
    if "due date" not in lower and "due  date" not in lower:
        return None
    for pat, code in ((ANC4_RE, "ANC4"), (ANC3_RE, "ANC3"), (ANC2_RE, "ANC2"), (ANC1_RE, "ANC1")):
        if pat.search(lower):
            return code
    return None


def main() -> int:
    if not SRC.exists():
        print(f"missing: {SRC.resolve()}")
        return 1
    shutil.copyfile(SRC, DST)
    wb = load_workbook(filename=str(DST))
    ws = wb["Question Master"]

    # Unmerge every range on the sheet — preserve the top-left cell's value but make every
    # cell writable. The source file uses merged ranges in column B (Form Name) which would
    # otherwise raise "MergedCell is read-only" when we try to write per-row values.
    for rng in list(ws.merged_cells.ranges):
        top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
        ws.unmerge_cells(str(rng))
        if top_left is not None:
            ws.cell(row=rng.min_row, column=rng.min_col).value = top_left

    # Locate the header row + relevant columns.
    header_row = -1
    cols: dict[str, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=30):
        for cell in row:
            key = header_key(cell.value)
            if key in ("faqid", "faquestionid", "qid"):
                header_row = cell.row
                break
        if header_row > 0:
            for cell in ws[header_row]:
                key = header_key(cell.value)
                if key in ("formcode", "form", "formname", "section", "sectioncode") and "form" not in cols:
                    cols["form"] = cell.column
                elif key in ("faqid", "faquestionid", "qid") and "qid" not in cols:
                    cols["qid"] = cell.column
                elif key in ("faquestionen", "questionen", "question", "faquestion") and "qen" not in cols:
                    cols["qen"] = cell.column
                elif key in ("faanswertype", "answertype", "type") and "atype" not in cols:
                    cols["atype"] = cell.column
            break

    if "qid" not in cols or "qen" not in cols or "atype" not in cols:
        print("Could not locate required columns (faQID, faQuestionEN, faAnswerType).")
        return 1

    # If a Form column wasn't already there, add one labelled "Form" right after faQID.
    if "form" not in cols:
        ws.insert_cols(idx=cols["qid"] + 1)
        cols["form"] = cols["qid"] + 1
        # shift cached column indices that are to the right of the inserted column
        for k in ("qen", "atype"):
            if cols[k] >= cols["form"]:
                cols[k] += 1
        ws.cell(row=header_row, column=cols["form"], value="Form")
        ws.cell(row=header_row, column=cols["form"]).fill = HIGHLIGHT

    # Walk rows and fill the Form column where it's blank.
    current: Optional[str] = None
    filled = 0
    untouched_meta = 0
    for r in range(header_row + 1, ws.max_row + 1):
        qid_val = ws.cell(row=r, column=cols["qid"]).value
        qen_val = ws.cell(row=r, column=cols["qen"]).value
        atype_val = ws.cell(row=r, column=cols["atype"]).value
        if is_blank(qid_val) and is_blank(qen_val):
            continue
        if is_blank(atype_val):
            # Geography / staff meta — these will be skipped by the parser.
            untouched_meta += 1
            continue

        existing = ws.cell(row=r, column=cols["form"]).value
        resolved: Optional[str] = None
        if not is_blank(existing):
            n = normalise(str(existing))
            resolved = FORM_CODE_ALIASES.get(n, n)
        if resolved is None:
            resolved = infer(str(qen_val))
        if resolved is not None:
            current = resolved
        if current is None:
            current = DEFAULT_FORM

        cell = ws.cell(row=r, column=cols["form"])
        if cell.value != current:
            cell.value = current
            cell.fill = HIGHLIGHT
            filled += 1

    wb.save(str(DST))
    print(f"wrote: {DST.resolve()}")
    print(f"  cells filled / changed: {filled}")
    print(f"  meta rows left untouched (no Form needed): {untouched_meta}")
    print(f"  Form column header at column {cols['form']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
