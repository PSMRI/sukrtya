"""Quick dump of SukrtyaQuestion.xlsx structure so we can confirm column names match the
FormExcelImportService parser. Not used at runtime — run manually with `python scripts/inspect_question_xlsx.py`.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def dump_sheet(ws, *, max_rows: int = 60, max_cols: int = 16) -> None:
    last_col = min(ws.max_column, max_cols)
    last_row = min(ws.max_row, max_rows)
    print(f"  dimensions: {ws.dimensions}  rows={ws.max_row}  cols={ws.max_column}")
    print(f"  showing first {last_row} rows x {last_col} cols")
    for r in range(1, last_row + 1):
        cells = []
        for c in range(1, last_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                cells.append("")
            else:
                s = str(v).replace("\n", " ").replace("\r", " ")
                if len(s) > 38:
                    s = s[:35] + "..."
                cells.append(s)
        print(f"  R{r:3d}: " + " | ".join(cells))
    print()


def main() -> int:
    path = Path("SukrtyaQuestion.xlsx")
    if not path.exists():
        print(f"missing: {path.resolve()}")
        return 1
    wb = load_workbook(filename=str(path), data_only=True, read_only=False)
    print(f"workbook: {path}")
    print(f"sheets:   {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n=== sheet: '{name}' ===")
        dump_sheet(ws)
    return 0


if __name__ == "__main__":
    sys.exit(main())
