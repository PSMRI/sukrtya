"""Dry-run that mirrors FormExcelImportService against SukrtyaQuestion.xlsx.

This is a *simulator* — it does not touch the DB. It applies the same parsing rules the
Java service uses so you can verify the form catalog + option sets that an upload would
produce. Run with: `python scripts/preview_form_import.py`.
"""
from __future__ import annotations

import re
from collections import OrderedDict
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


FORM_CODE_ALIASES = {
    "BASIC": "BASIC",
    "BASICDETAILS": "BASIC",
    "BASICDETAIL": "BASIC",
    "PREGNANTWOMAN": "BASIC",
    "ANC1": "ANC1",
    "ANC2": "ANC2",
    "ANC3": "ANC3",
    "ANC4": "ANC4",
    "FORM1": "BASIC",
    "1STANC": "ANC1",
    "2NDANC": "ANC2",
    "3RDANC": "ANC3",
    "4THANC": "ANC4",
    "FIRSTANC": "ANC1",
    "SECONDANC": "ANC2",
    "THIRDANC": "ANC3",
    "FOURTHANC": "ANC4",
}

ANC1_RE = re.compile(r"\b(1st\s*anc|first\s*anc|anc\s*1\b|of\s*1\s*anc)\b", re.I)
ANC2_RE = re.compile(r"\b(2nd\s*anc|second\s*anc|anc\s*2\b|of\s*2\s*anc)\b", re.I)
ANC3_RE = re.compile(r"\b(3rd\s*anc|third\s*anc|anc\s*3\b|of\s*3\s*anc)\b", re.I)
ANC4_RE = re.compile(r"\b(4th\s*anc|fourth\s*anc|anc\s*4\b|of\s*4\s*anc)\b", re.I)

DATE_OFFSET_RE = re.compile(r"\s*([A-Za-z][A-Za-z0-9_ ]+?)\s*\+\s*(\d+)\s*Days?\s*", re.I)

DEFAULT_FORM_CODE = "BASIC"


def header_key(s: object) -> str:
    if s is None:
        return ""
    return str(s).strip().lower().replace(" ", "")


def is_blank(s: object) -> bool:
    if s is None:
        return True
    t = str(s).strip()
    return t == "" or t.lower() == "null" or t == "-"


def parse_int(s: object) -> Optional[int]:
    if s is None:
        return None
    t = str(s).strip()
    if not t or t.lower() == "null":
        return None
    try:
        return int(t)
    except ValueError:
        try:
            return int(float(t))
        except ValueError:
            return None


def parse_answer_type(s: object) -> Optional[str]:
    if s is None:
        return None
    t = str(s).strip().lower().replace(" ", "").replace("_", "")
    return {
        "text": "TEXT", "string": "TEXT",
        "numeric": "NUMERIC", "number": "NUMERIC", "int": "NUMERIC", "integer": "NUMERIC",
        "decimal": "NUMERIC", "float": "NUMERIC", "double": "NUMERIC",
        "date": "DATE",
        "singlechoice": "SINGLE_CHOICE", "single": "SINGLE_CHOICE", "radio": "SINGLE_CHOICE",
        "dropdown": "SINGLE_CHOICE", "select": "SINGLE_CHOICE",
        "multichoice": "MULTI_CHOICE", "multi": "MULTI_CHOICE", "multiple": "MULTI_CHOICE",
        "checkbox": "MULTI_CHOICE", "checkboxes": "MULTI_CHOICE",
    }.get(t)


def parse_bool(s: object) -> bool:
    if s is None:
        return False
    t = str(s).strip().lower()
    if t in ("", "-", "null"):
        return False
    return t in ("1", "true", "yes", "y")


def parse_option_ids(raw: object) -> list[int]:
    if raw is None or is_blank(raw):
        return []
    tokens = re.split(r"[/,;\s]+", str(raw).strip())
    seen: list[int] = []
    for t in tokens:
        v = parse_int(t)
        if v is not None and v not in seen:
            seen.append(v)
    return seen


def normalise_form_code(raw: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", raw.strip().upper())


def simplify_label(s: Optional[str]) -> str:
    if not s:
        return ""
    return re.sub(r"[^a-z0-9]", "", s.lower())


def strip_date_suffix(s: str) -> str:
    return re.sub(r"\bdate\b", "", s, flags=re.I).strip()


def infer_section(label: Optional[str]) -> Optional[str]:
    if not label:
        return None
    lower = label.lower()
    if "due date" not in lower and "due  date" not in lower:
        return None
    for pat, code in ((ANC4_RE, "ANC4"), (ANC3_RE, "ANC3"), (ANC2_RE, "ANC2"), (ANC1_RE, "ANC1")):
        if pat.search(lower):
            return code
    return None


def build_option_set_code(sorted_ids: list[int]) -> str:
    code = "OPT_" + "_".join(str(x) for x in sorted_ids)
    if len(code) <= 50:
        return code
    short = f"OPT_{sorted_ids[0]}_TO_{sorted_ids[-1]}_N{len(sorted_ids)}"
    return short[:50]


def main() -> int:
    path = Path("SukrtyaQuestion.xlsx")
    wb = load_workbook(filename=str(path), data_only=True)

    # ------- OptionMaster -------
    options: dict[int, tuple[str, Optional[str]]] = {}
    opt = wb["OptionMaster"]
    rows = list(opt.iter_rows(values_only=True))
    header = [header_key(c) for c in rows[0]]
    try:
        idx_id = next(i for i, h in enumerate(header) if h in ("optionid", "id"))
    except StopIteration:
        print("OptionMaster header missing OptionID column")
        return 1
    idx_en = next((i for i, h in enumerate(header) if h in ("optionnameen", "labelen", "label", "name", "nameen")), None)
    idx_hi = next((i for i, h in enumerate(header) if h in ("optionnamereg", "labelhi", "labelregional", "namereg", "namehi", "regional")), None)
    if idx_en is None:
        print("OptionMaster header missing label column")
        return 1
    for r in rows[1:]:
        oid = parse_int(r[idx_id])
        if oid is None:
            continue
        en = r[idx_en]
        if en is None or str(en).strip() == "":
            continue
        hi = r[idx_hi] if (idx_hi is not None and idx_hi < len(r)) else None
        options[oid] = (str(en).strip(), str(hi).strip() if hi not in (None, "") else None)

    # ------- Question Master -------
    qm = wb["Question Master"]
    rows = list(qm.iter_rows(values_only=True))
    # find header row
    header_row_idx = -1
    for i, r in enumerate(rows[:30]):
        for c in r:
            if header_key(c) in ("faqid", "faquestionid", "qid"):
                header_row_idx = i
                break
        if header_row_idx >= 0:
            break
    if header_row_idx < 0:
        print("Question Master header row not found")
        return 1
    header = [header_key(c) for c in rows[header_row_idx]]

    def col(*candidates: str) -> Optional[int]:
        for i, h in enumerate(header):
            if h in candidates:
                return i
        return None

    c_form = col("formcode", "form", "formname", "section", "sectioncode")
    c_qid = col("faqid", "faquestionid", "qid")
    c_qen = col("faquestionen", "questionen", "question", "faquestion")
    c_qhi = col("faquestionhi", "questionhi", "faquestionreg", "questionreg")
    c_ans = col("faanswer", "defaultanswer", "answer")
    c_skipa = col("skipanswer", "skipifvalue", "skipif")
    c_skipto = col("skiptoquestion", "skipto", "skiptoqid")
    c_type = col("faanswertype", "answertype", "type")
    c_mand = col("ismandatory", "ismandate", "mandatory", "required")
    c_max = col("maxvalue", "max")
    c_min = col("minvalue", "min")
    c_rem = col("remarks", "remark", "note", "notes", "formula")

    grouped: OrderedDict[str, list[dict]] = OrderedDict()
    current_section: Optional[str] = None
    skipped = 0

    def cell(r: tuple, i: Optional[int]) -> Optional[object]:
        if i is None or i >= len(r):
            return None
        return r[i]

    for ri, r in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
        qid_raw = cell(r, c_qid)
        q_en = cell(r, c_qen)
        atype_raw = cell(r, c_type)
        if is_blank(qid_raw) and is_blank(q_en):
            continue
        if is_blank(atype_raw):
            # geography/meta rows
            continue
        qid = parse_int(qid_raw)
        if qid is None:
            skipped += 1
            continue
        if is_blank(q_en):
            skipped += 1
            continue
        atype = parse_answer_type(atype_raw)
        if atype is None:
            skipped += 1
            continue

        form_raw = cell(r, c_form)
        resolved: Optional[str] = None
        if not is_blank(form_raw):
            n = normalise_form_code(str(form_raw))
            resolved = FORM_CODE_ALIASES.get(n, n)
        if resolved is None:
            resolved = infer_section(str(q_en))
        if resolved is not None:
            current_section = resolved
        if current_section is None:
            current_section = DEFAULT_FORM_CODE

        mand = parse_bool(cell(r, c_mand))
        fa = cell(r, c_ans)
        opt_ids = parse_option_ids(fa) if atype in ("SINGLE_CHOICE", "MULTI_CHOICE") else []

        grouped.setdefault(current_section, []).append(dict(
            row=ri,
            faQid=qid,
            label=str(q_en).strip(),
            type=atype,
            mandatory=mand,
            optionIds=opt_ids,
            faAnswer=fa,
            remarks=cell(r, c_rem),
            skipAnswer=cell(r, c_skipa),
            skipTo=cell(r, c_skipto),
            minValue=cell(r, c_min),
            maxValue=cell(r, c_max),
        ))

    # Build global label index for computed-field resolution
    global_labels: dict[str, str] = {}
    for code, drafts in grouped.items():
        for d in drafts:
            global_labels[simplify_label(d["label"])] = f"{code}_{d['faQid']}"

    # Resolve computed fields
    computed = []
    for code, drafts in grouped.items():
        for d in drafts:
            rem = d["remarks"]
            if rem is None:
                continue
            m = DATE_OFFSET_RE.fullmatch(str(rem).strip())
            if not m:
                continue
            src_label = simplify_label(m.group(1))
            src_code = global_labels.get(src_label) or global_labels.get(simplify_label(strip_date_suffix(m.group(1))))
            computed.append(dict(
                code=f"{code}_{d['faQid']}",
                label=d["label"],
                source_raw=m.group(1).strip(),
                source_code=src_code,
                offset_days=int(m.group(2)),
            ))

    # Collect option sets
    set_codes: dict[tuple[int, ...], str] = {}
    for code, drafts in grouped.items():
        for d in drafts:
            ids = d["optionIds"]
            if not ids:
                continue
            key = tuple(sorted(set(ids)))
            if key not in set_codes:
                set_codes[key] = build_option_set_code(list(key))

    # ------- Print summary -------
    print("=" * 78)
    print(f"DRY-RUN PREVIEW for {path}")
    print("=" * 78)

    print(f"\nOptionMaster: loaded {len(options)} option labels.")
    missing_labels_per_set: dict[tuple[int, ...], list[int]] = {}
    for ids, set_code in set_codes.items():
        miss = [i for i in ids if i not in options]
        if miss:
            missing_labels_per_set[ids] = miss

    print(f"\nForms ({len(grouped)}):")
    for code, drafts in grouped.items():
        print(f"  [{code}] {len(drafts)} questions")
        for d in drafts:
            opt_summary = ""
            if d["optionIds"]:
                key = tuple(sorted(set(d["optionIds"])))
                set_code = set_codes[key]
                opt_summary = f"  -> {set_code}"
            print(f"    {code}_{d['faQid']:>2}  {d['type']:<14} mand={int(d['mandatory'])}  "
                  f"{d['label'][:50]}{opt_summary}")

    print(f"\nOption sets to be auto-created ({len(set_codes)}):")
    for ids, set_code in set_codes.items():
        labels = [f"{i}={options[i][0]}" if i in options else f"{i}=<missing>" for i in ids]
        print(f"  {set_code}")
        for line in labels:
            print(f"      {line}")

    if missing_labels_per_set:
        print(f"\nWARNING: option IDs referenced in Question Master but missing from OptionMaster:")
        for ids, miss in missing_labels_per_set.items():
            print(f"  set {set_codes[ids]} -> missing ids {miss}")

    print(f"\nComputed (DATE_OFFSET) fields ({len(computed)}):")
    for c in computed:
        src = c["source_code"] or "<UNRESOLVED>"
        print(f"  {c['code']:>12}  {c['label']:<28} = {src} + {c['offset_days']} days  (raw remark: '{c['source_raw']}')")
    unresolved = [c for c in computed if not c["source_code"]]
    if unresolved:
        print(f"\nWARNING: {len(unresolved)} computed field(s) could not resolve their source.")

    print(f"\nTotals: questions={sum(len(v) for v in grouped.values())}  "
          f"skipped={skipped}  optionSets={len(set_codes)}  computed={len(computed)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
